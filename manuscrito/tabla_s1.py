#!/usr/bin/env python3
"""Supplementary Table S1 (English): guideline value database of Kenti Water Quality 0.2,
built from repo/datos/niveles_guia.csv. Output: Flores_Kenti_Water_Quality_JWET_Supplementary.docx"""
import csv
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).parent
rows = list(csv.DictReader(open(HERE.parent / "repo" / "datos" / "niveles_guia.csv", encoding="utf-8-sig")))

NORMA = {
    "caa982": ("Argentine Food Code, Art. 982", "[4]"), "who2022": ("WHO Guidelines for Drinking-water Quality (2022)", "[1]"),
    "ue2020": ("Directive (EU) 2020/2184, Annex I", "[2]"), "epa": ("US EPA primary and secondary drinking water regulations", "[3]"),
    "d831_t1": ("Decree 831/93, Annex II, Table 1: sources for drinking water", "[5]"), "d831_t2": ("Decree 831/93, Annex II, Table 2: aquatic life (fresh water)", "[5]"),
    "d831_t5": ("Decree 831/93, Annex II, Table 5: irrigation", "[5]"), "d831_t6": ("Decree 831/93, Annex II, Table 6: livestock water", "[5]"),
    "l24585_t1": ("Law 24.585, Annex IV, Table 1", "[6]"), "l24585_t2": ("Law 24.585, Annex IV, Table 2", "[6]"),
    "l24585_t5": ("Law 24.585, Annex IV, Table 5", "[6]"), "l24585_t6": ("Law 24.585, Annex IV, Table 6", "[6]"),
    "fao29": ("FAO Irrigation and Drainage Paper 29 Rev. 1", "[12]"), "ccme_ganado": ("CCME livestock water guidelines", "[10]"),
    "anzg_ganado": ("ANZECC/ARMCANZ (2000) livestock drinking water", "[11]"),
    "jp_suido": ("Japan: drinking water quality standards (水質基準, 52 items)", "[7]"),
    "jp_eqs_salud": ("Japan: environmental quality standards, human health (健康項目)", "[8]"),
    "jp_rio": ("Japan: environmental quality standards for rivers, living environment (生活環境項目, classes AA–E)", "[8]"),
    "jp_biota": ("Japan: environmental quality standard for aquatic organisms in rivers", "[8]"),
    "jp_ind": ("Japan: river classes suitable for industrial water (classes C–E)", "[8]"),
    "jp_riego": ("Japan: agricultural (paddy rice) water standard (農業（水稲）用水基準)", "[9]"),
}
PARAM = {"Aluminio": "Aluminium", "Amonio": "Ammonium", "Antimonio": "Antimony", "Arsénico": "Arsenic", "Bacterias aerobias mesófilas": "Heterotrophic plate count",
    "Bario": "Barium", "Berilio": "Beryllium", "Bicarbonato": "Bicarbonate", "Boro": "Boron", "Bromato": "Bromate", "Cadmio": "Cadmium", "Calcio": "Calcium",
    "Carbono orgánico total": "Total organic carbon", "Cianuro": "Cyanide", "Cinc": "Zinc", "Cloro residual": "Residual chlorine", "Cloruro": "Chloride",
    "Cobalto": "Cobalt", "Cobre": "Copper", "Coliformes totales": "Total coliforms", "Color": "Colour", "Conductividad eléctrica": "Electrical conductivity",
    "Cromo hexavalente": "Hexavalent chromium", "Cromo total": "Total chromium", "DBO₅": "BOD₅", "DQO al permanganato (COD-Mn)": "Permanganate COD (COD-Mn)",
    "Dureza total": "Total hardness", "Escherichia coli": "Escherichia coli", "Fluoruro": "Fluoride", "Hierro": "Iron", "Litio": "Lithium",
    "Manganeso": "Manganese", "Mercurio": "Mercury", "Molibdeno": "Molybdenum", "Nitrato": "Nitrate", "Nitrito": "Nitrite", "Nitrógeno total": "Total nitrogen",
    "Níquel": "Nickel", "Oxígeno disuelto": "Dissolved oxygen", "PFOS + PFOA (suma)": "PFOS + PFOA (sum)", "Plata": "Silver", "Plomo": "Lead",
    "Pseudomonas aeruginosa": "Pseudomonas aeruginosa", "RAS (relación de adsorción de sodio)": "SAR (sodium adsorption ratio)", "Selenio": "Selenium",
    "Sodio": "Sodium", "Sulfato": "Sulfate", "Sólidos disueltos totales": "Total dissolved solids", "Sólidos suspendidos totales": "Total suspended solids",
    "Talio": "Thallium", "Turbiedad": "Turbidity", "Uranio": "Uranium", "Vanadio": "Vanadium", "pH": "pH"}
TIPO = {"salud": "Health", "estetico": "Acceptability", "indicador": "Indicator", "secundario": "Secondary", "accion": "Action level",
        "operativo": "Operational", "grado": "Degree of restriction"}
UNID = {"UFC/100 mL": "CFU/100 mL", "UFC/mL": "CFU/mL", "unid. pH": "pH units", "mg/L CaCO₃": "mg/L as CaCO₃"}
COND = {"": "",
    "0,005 o 0,1 mg/L según pH, Ca y COD": "0.005 or 0.1 mg/L depending on pH, Ca and DOC",
    "Por especie: de 2000 a 5000 mg/L": "By species: 2000 to 5000 mg/L",
    "Por especie: ovinos 0,4; bovinos 1; porcinos y aves 5": "By species: sheep 0.4; cattle 1; pigs and poultry 5",
    "Por especie: ovinos 0,5; bovinos 1; porcinos y aves 5": "By species: sheep 0.5; cattle 1; pigs and poultry 5",
    "Según clase de agua industrial 1–3": "By industrial water class 1–3", "Según clase de río AA–E": "By river class AA–E",
    "Según temperatura media anual: de 1,7 a 0,8 mg/L": "By mean annual temperature: 1.7 to 0.8 mg/L"}
NOTA = {"": "",
    "0 en 100 ml.": "0 in 100 mL.",
    "0,10 mg/l a la salida de la planta.": "0.10 mg/L at the treatment plant outlet.",
    "0,3 mS/cm.": "0.3 mS/cm.",
    "1,5 y 8,5 meq/l; sólo aspersión sobre follaje.": "1.5 and 8.5 meq/L; overhead sprinkling only.",
    "1.000–2.000 mg/l: posibles efectos; > 2.000 mg/l: problemas crónicos o agudos.": "1000–2000 mg/L: possible effects; > 2000 mg/L: chronic or acute problems.",
    "1–2 mg/l; el menor si el alimento también aporta flúor.": "1–2 mg/L; the lower value if feed also contains fluoride.",
    "2 grados de turbiedad.": "2 turbidity units.",
    "2 µg/l para protección de la vida acuática incluido fito y zooplancton; 20 µg/l sólo para peces (nota 7).": "2 µg/L to protect aquatic life including phyto- and zooplankton; 20 µg/L for fish only (note 7).",
    "2,4 mg/l desde la Res. Conj. 33/2023 (antes 0,5 mg/l).": "2.4 mg/L since Joint Resolution 33/2023 (previously 0.5 mg/L).",
    "25 µg/l desde el 12/01/2036; hasta entonces, 50 µg/l.": "25 µg/L from 12 January 2036; 50 µg/L until then.",
    "30 µg/l en zonas con geología que lo justifique.": "30 µg/L where local geology justifies it.",
    "4 y 10 meq/l, riego superficial. En aspersión: > 3 meq/l (106 mg/l) ya restringe.": "4 and 10 meq/L, surface irrigation. With sprinklers, > 3 meq/L (106 mg/L) already restricts use.",
    "5 grados de color.": "5 colour units.",
    "5 µg/l con pH < 6,5, Ca²⁺ < 4 mg/l y COD < 2 mg/l; 100 µg/l con pH ≥ 6,5, Ca²⁺ ≥ 4 mg/l y COD ≥ 2 mg/l (nota 2 de la tabla 2).": "5 µg/L with pH < 6.5, Ca²⁺ < 4 mg/L and DOC < 2 mg/L; 100 µg/L with pH ≥ 6.5, Ca²⁺ ≥ 4 mg/L and DOC ≥ 2 mg/L (note 2 of Table 2).",
    "5 µg/l desde el 12/01/2036; hasta entonces, 10 µg/l.": "5 µg/L from 12 January 2036; 10 µg/L until then.",
    "A 20 °C.": "At 20 °C.",
    "Además: [nitrato]/50 + [nitrito]/3 ≤ 1 (ver «Iones y derivados»).": "Also: [nitrate]/50 + [nitrite]/3 ≤ 1 (see “Ions and derived quantities”).",
    "Aluminio residual.": "Residual aluminium.",
    "Ausencia en 100 ml.": "Absent in 100 mL.",
    "Bacterias generales: hasta 100 colonias en 1 mL.": "General bacteria: up to 100 colonies in 1 mL.",
    "Calcio, magnesio, etc. (dureza).": "Calcium, magnesium, etc. (hardness).",
    "Cinc total, promedio anual, igual en las clases 生物A, 生物特A, 生物B y 生物特B.": "Total zinc, annual mean, same for classes 生物A, 生物特A, 生物B and 生物特B.",
    "Como cianuro libre.": "As free cyanide.",
    "Como cianuro libre. Si el laboratorio informó cianuro total, la comparación es conservadora.": "As free cyanide. If the laboratory reported total cyanide, the comparison is conservative.",
    "Considerar el volumen consumido y el aporte de otras fuentes.": "Consider the volume consumed and intake from other sources.",
    "Cromo total.": "Total chromium.",
    "DQO al permanganato (método japonés), no al dicromato.": "Permanganate COD (Japanese method), not dichromate COD.",
    "Depende de la dureza (mg/l CaCO₃): 0–60, 60–120, 120–180 y > 180 (notas de la tabla 2).": "Depends on hardness (mg/L as CaCO₃): 0–60, 60–120, 120–180 and > 180 (notes of Table 2).",
    "El decreto no aclara la forma; el valor es el de la EPA (fuente A de la tabla), expresado como N.": "The decree does not state the form; the value is the EPA value (source A of the table), expressed as N.",
    "El decreto no aclara la forma; el valor es el de la EPA (fuente B), expresado como N.": "The decree does not state the form; the value is the EPA value (source B), expressed as N.",
    "El límite depende de la temperatura media y máxima del año del lugar (tabla del art. 982): de 1,7 mg/l (10–12 °C) a 0,8 mg/l (26,3–32,6 °C). Se carga por sitio.": "The limit depends on the site’s annual mean and maximum temperature (table in Art. 982): from 1.7 mg/L (10–12 °C) to 0.8 mg/L (26.3–32.6 °C). Entered per site.",
    "El nivel secundario es 2,0 mg/l.": "The secondary level is 2.0 mg/L.",
    "En la norma se evalúa como valor del 90 % de los muestreos del año; con muestras sueltas la comparación es orientativa.": "The standard is assessed on the 90% value of the year’s samples; with individual samples the comparison is indicative.",
    "Gusto y manchas desde unos 0,3 mg/l (cap. 10).": "Taste and staining from about 0.3 mg/L (Ch. 10).",
    "Hasta 1.500 mg/l puede tolerarse si el alimento tiene poco nitrato.": "Up to 1500 mg/L may be tolerated if feed is low in nitrate.",
    "Hasta 5 mg/l si el arsénico no se agrega como aditivo del alimento.": "Up to 5 mg/L if arsenic is not added as a feed additive.",
    "Hasta 500 UFC/ml.": "Up to 500 CFU/mL.",
    "Ion cianuro y cloruro de cianógeno, como cianuro.": "Cyanide ion and cyanogen chloride, as cyanide.",
    "La Res. Conj. 33/2023 admite hasta 0,4 mg/l en zonas con alto contenido natural de manganeso.": "Joint Resolution 33/2023 allows up to 0.4 mg/L in areas with naturally high manganese.",
    "Materia orgánica como carbono orgánico total.": "Organic matter as total organic carbon.",
    "Mercurio inorgánico.": "Inorganic mercury.",
    "Mercurio total.": "Total mercury.",
    "Mínimo exigible en agua de red desinfectada con cloro; no corresponde en agua sin tratar.": "Minimum required in chlorinated supply water; not applicable to untreated water.",
    "Nitrato + nitrito, como N. (suma nitrato + nitrito)": "Nitrate + nitrite, as N (sum of nitrate and nitrite).",
    "Nitrito como N.": "Nitrite as N.",
    "Nitrito solo, como N.": "Nitrite alone, as N.",
    "Nitrógeno como nitrato (NO₃-N); afecta cultivos sensibles.": "Nitrate nitrogen (NO₃-N); affects sensitive crops.",
    "Nitrógeno total.": "Total nitrogen.",
    "Nivel de acción (tratamiento). El nivel secundario es 1,0 mg/l.": "Action level (treatment). The secondary level is 1.0 mg/L.",
    "Nivel de acción. La regla LCRI (2024) lo baja a 0,010 mg/l desde noviembre de 2027.": "Action level. The LCRI rule (2024) lowers it to 0.010 mg/L from November 2027.",
    "No detectable en 100 ml.": "Not detectable in 100 mL.",
    "No detectable.": "Not detectable.",
    "No se aplica en el mar.": "Not applicable to sea water.",
    "Para mercurio inorgánico.": "For inorganic mercury.",
    "Para regiones con alto contenido natural de arsénico, el CAA previó un plazo de adecuación desde 0,05 mg/l, prorrogado hasta completar los estudios de hidroarsenicismo. Verificar la situación de la jurisdicción.": "For regions with naturally high arsenic, the Food Code set a transition period starting from 0.05 mg/L, extended until hydroarsenicism studies are completed. Check the situation in each jurisdiction.",
    "Por debajo de este valor puede manchar ropa y sanitarios.": "Below this value it may stain laundry and sanitary ware.",
    "Por encima de unos 1.000 mg/l el agua se vuelve cada vez menos aceptable al gusto (cap. 10). No es un valor de salud.": "Above about 1000 mg/L water becomes increasingly unpalatable (Ch. 10). Not a health-based value.",
    "Por especie: ovinos 0,4; bovinos 1; porcinos y aves 5 mg/l.": "By species: sheep 0.4; cattle 1; pigs and poultry 5 mg/L.",
    "Por especie: ovinos 0,5; bovinos 1; porcinos y aves 5 mg/l.": "By species: sheep 0.5; cattle 1; pigs and poultry 5 mg/L.",
    "Provisional (analítica y tratamiento).": "Provisional (analytical and treatment achievability).",
    "Provisional.": "Provisional.",
    "Provisional. Manganeso total; también afecta color y gusto.": "Provisional. Total manganese; also affects colour and taste.",
    "Provisional: se fijó por lo que se puede medir y tratar, no sólo por salud.": "Provisional: set on what can be measured and treated, not on health alone.",
    "Rango normal.": "Normal range.",
    "Rango secundario 0,05–0,2 mg/l.": "Secondary range 0.05–0.2 mg/L.",
    "Residuo por evaporación (蒸発残留物): se compara con los sólidos disueltos totales.": "Evaporation residue (蒸発残留物): compared with total dissolved solids.",
    "Salinidad: < 0,7 dS/m sin restricción; 0,7–3,0 ligera a moderada; > 3,0 severa.": "Salinity: < 0.7 dS/m no restriction; 0.7–3.0 slight to moderate; > 3.0 severe.",
    "Sin efectos adversos hasta: bovinos de carne 4.000; de leche 2.500; ovinos 5.000; equinos y porcinos 4.000; aves 2.000 mg/l (tabla 4.3.1).": "No adverse effects up to: beef cattle 4000; dairy cattle 2500; sheep 5000; horses and pigs 4000; poultry 2000 mg/L (Table 4.3.1).",
    "Suma de PFOS y PFOA: 50 ng/L. Pasó de objetivo a norma el 1/4/2026.": "Sum of PFOS and PFOA: 50 ng/L. Changed from a target to a standard on 1 April 2026.",
    "Toxicidad por sodio en riego superficial (RAS). La infiltración (RAS con CE) va en «Iones y derivados».": "Sodium toxicity in surface irrigation (SAR). Infiltration (SAR with EC) is assessed under “Ions and derived quantities”.",
    "Umbral de gusto aproximado (cap. 10).": "Approximate taste threshold (Ch. 10).",
    "Umbral de gusto aproximado: 200–300 mg/l (cap. 10).": "Approximate taste threshold: 200–300 mg/L (Ch. 10).",
    "Umbral de gusto aproximado: 250–1.000 mg/l según el catión (cap. 10).": "Approximate taste threshold: 250–1000 mg/L depending on the cation (Ch. 10).",
    "Valor impreso en el decreto: 50 µg/l. La fuente que cita (CCREM 1987) fija 50 mg/l para ganado, mil veces más: probable error de transcripción. Kenti compara con el valor impreso.": "Value printed in the decree: 50 µg/L. The source it cites (CCREM 1987) sets 50 mg/L for livestock, a thousand times higher: probable transcription error. Kenti compares against the printed value.",
    "«Amonio (total)» 1.370 µg/l. El decreto no aclara la forma; se compara como NH₄⁺.": "“Ammonium (total)” 1370 µg/L. The decree does not state the form; compared as NH₄⁺.",
    "«No detectable» con límite de cuantificación 0,1 mg/L: se considera que supera si se cuantifica.": "“Not detectable” with a quantification limit of 0.1 mg/L: considered exceeded if quantified.",
}
FORMA = {"": "", "N": "as N", "NH4": "as NH₄⁺", "NO2": "as NO₂⁻", "NO3": "as NO₃⁻"}

def num(x):
    return x.replace(",", ".") if x else ""

def valor(r):
    lo, hi, l1, l2 = num(r["minimo"]), num(r["maximo"]), num(r["umbral_ligero"]), num(r["umbral_severo"])
    if l1 or l2:
        return f"slight {l1}; severe {l2}"
    if lo and hi:
        return f"{lo}–{hi}"
    if hi:
        return f"≤ {hi}"
    if lo:
        return f"≥ {lo}"
    return "—"

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.8))
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(8)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "MS Mincho")

def par(text, size=11, bold=False, after=6):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    return p

par("Supplementary Materials", 14, True)
par("Kenti Water Quality: an open-source tool for assessing water analyses against multiple regulatory frameworks with incomplete data", 11, False)
par("Carlos María Flores", 11, False, 12)
par("Table S1. Guideline value database of Kenti Water Quality 0.2 (373 values in 21 frameworks). Values are in the unit shown and in the chemical form given in the “Form” column. "
    "Type: health-based, acceptability, indicator, secondary, action level, operational, or degree of restriction (FAO). Conditional limits are resolved per sample or site as described in the Materials and Methods; "
    "the “Condition” column states the dependence. Reference numbers correspond to the reference list of the main text. "
    "The same table is distributed in machine-readable form (datos/niveles_guia.csv) in the software repository.", 9, False, 8)

cols = ["Parameter", "Value", "Unit", "Form", "Type", "Condition", "Note"]
widths = [Cm(3.6), Cm(2.6), Cm(2.0), Cm(1.6), Cm(2.2), Cm(4.0), Cm(10.1)]
actual = None; tabla = None

def encabezado_repetido(row):
    trPr = row._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)

for r in rows:
    if r["norma_id"] != actual:
        actual = r["norma_id"]; nombre, ref = NORMA[actual]
        p = par(f"{nombre} {ref}", 9, True, 3); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
        tabla = doc.add_table(rows=1, cols=len(cols)); tabla.style = "Table Grid"; tabla.autofit = False
        for i, c in enumerate(cols):
            cell = tabla.rows[0].cells[i]; cell.width = widths[i]; cell.text = ""
            run = cell.paragraphs[0].add_run(c); run.bold = True; run.font.size = Pt(8)
        for i, w in enumerate(widths):
            tabla.columns[i].width = w
        tblW = tabla._tbl.tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); tabla._tbl.tblPr.append(tblW)
        tblW.set(qn("w:w"), str(sum(int(w.twips) for w in widths))); tblW.set(qn("w:type"), "dxa")
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tabla._tbl.tblPr.append(lay)
        encabezado_repetido(tabla.rows[0])
    fila = tabla.add_row().cells
    vals = [PARAM[r["parametro"]], valor(r), UNID.get(r["unidad"], r["unidad"]), FORMA[r["expresado_como"]], TIPO[r["tipo_de_valor"]],
            COND[r["condicion"]] if r["condicion"] in COND else r["condicion"].replace("Según dureza", "By hardness").replace(" para ", " for ").replace(" y >", " and >").replace("CaCO3", "as CaCO₃"),
            NOTA[r["nota"]]]
    for i, v in enumerate(vals):
        fila[i].width = widths[i]; fila[i].text = v
        for p in fila[i].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            if i == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER

out = HERE / "Flores_Kenti_Water_Quality_JWET_Supplementary.docx"
doc.save(out)
print("Escrito", out, len(rows), "valores")
